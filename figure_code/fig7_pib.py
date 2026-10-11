"""
Figure 7 -- application to the 2 wt% polyisobutylene solution of King & McKinley (Sec. VIII of the manuscript).

Data: N. King, figshare (2026), doi:10.6084/m9.figshare.32005782 (CC BY 4.0), file PIB_freqsweep_flowcurve.xlsx,
      sheets 'Frequency - 1' (omega, G', G'') and 'Sweep - 1' (gdot, eta, Psi1).

Steps
  (A) Linear relaxation spectrum: non-negative least squares (NNLS) fits of G' and G'' - eta_s omega with relaxation times
      fixed on logarithmic grids (5, 7, 7 and 9 modes), and a chain-like spectrum tau_p = tau_1 p^-k with equal moduli
      (nonlinear least squares on ln G', ln G''). Moments eta0, <tau>, <tau^2> of each.
  (B) Composite from the measured flow curve (Eqs. 1-2), with and without subtraction of the solvent (eta_s = 18 mPa s).
  (C) Three multimode nonlinear models sharing the 7-mode spectrum -- linear PTT, exponential PTT, Giesekus -- with a
      mode-dependent parameter p_k = p_1 (tau_k/tau_max)^m fitted to ln eta (gdot <= 63 1/s) and ln Psi1
      (0.3 <= gdot <= 63 1/s); their true planar viscosities. (D) The single-mode Rolie-Poly model of King & McKinley.

Run:  python3 -I fig7_pib.py /path/to/PIB_freqsweep_flowcurve.xlsx    (about 5-10 min)
      (-I isolates the interpreter from the current directory, a precaution when reading a downloaded file)
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))     # needed with python3 -I
import numpy as np
from scipy.optimize import nnls, least_squares, minimize
from common import (lptt_shear, lptt_ext, eptt_shear, eptt_ext, giesekus_shear, giesekus_ext, rolie_poly, steady_state,
                    L_shear, L_ext, setup_matplotlib, panel_label, save_json, load_json, PALETTE, INK, MUTE, FIGS, Checks)

ETA_S = 0.018                    # solvent viscosity, Pa s
STABLE = 70.0                    # flow curve used up to the highest stable rate (63 1/s); beyond it the flow is unstable


def read_data(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    F = np.array([r[:4] for r in wb['Frequency - 1'].iter_rows(min_row=4, values_only=True) if r[0] is not None], float)[::-1]
    S = np.array([r[2:6] for r in wb['Sweep - 1'].iter_rows(min_row=4, values_only=True) if r[0] is not None], float)
    return dict(w=F[:, 1], Gp=F[:, 2], Gpp=F[:, 3], gd=S[:, 1], eta=S[:, 2], Psi1=S[:, 3])


def maxwell_moduli(tk, hk, w):
    """G', G'' of Maxwell modes with relaxation times tk and viscosities hk (moduli hk/tk)."""
    x = np.outer(w, tk)
    return (x**2 / (1 + x**2)) @ (hk / tk), (x / (1 + x**2)) @ (hk / tk)


def moments(tk, hk):
    e0 = hk.sum()
    return e0, (hk * tk).sum() / e0, (hk * tk**2).sum() / e0


def spectra(D):
    w, Gp, Gpp = D['w'], D['Gp'], D['Gpp']
    est = {}
    for name, tk in [('5 modes', np.logspace(-2.3, 1.0, 5)), ('7 modes', np.logspace(-2.3, 1.0, 7)),
                     ('7 modes, to 30 s', np.logspace(-2.3, 1.5, 7)), ('9 modes', np.logspace(-2.5, 1.2, 9))]:
        x = np.outer(w, tk)
        A = np.vstack([(x**2 / (1 + x**2)) / Gp[:, None], (x / (1 + x**2)) / (Gpp - ETA_S * w)[:, None]])
        g, res = nnls(A, np.ones(2 * len(w)))           # moduli g_k; relative residuals of G' and G'' - eta_s w
        hk = g * tk
        e0, t1, t2 = moments(tk, hk)
        est[name] = dict(tk=tk.tolist(), hk=hk.tolist(), rms=res / np.sqrt(2 * len(w)), eta0=e0 + ETA_S, t1=t1, t2=t2)

    def model(p, P=200):
        G, tau1, k = np.exp(p[0]), np.exp(p[1]), p[2]
        tk = tau1 * np.arange(1, P + 1.0)**-k
        return tk, G * tk

    def res(p):
        tk, hk = model(p)
        a, b = maxwell_moduli(tk, hk, w)
        return np.concatenate([np.log(a / Gp), np.log((b + ETA_S * w) / Gpp)])
    r = least_squares(res, [np.log(3.0), np.log(3.0), 1.8])
    cov = np.linalg.inv(r.jac.T @ r.jac) * np.sum(r.fun**2) / (len(r.fun) - 3)
    tk, hk = model(r.x)
    e0, t1, t2 = moments(tk, hk)
    est['chain'] = dict(tau1=float(np.exp(r.x[1])), k=r.x[2], dk=np.sqrt(cov[2, 2]), rms=np.sqrt(np.mean(r.fun**2)),
                        eta0=e0 + ETA_S, t1=t1, t2=t2)
    return est


def composite(D):
    """Composite from the measured flow curve; Psi1 below ~0.25 1/s is at the transducer noise floor and is not used."""
    m = (D['gd'] < STABLE) & (D['gd'] > 0.22)
    g, eta, Psi1 = D['gd'][m], D['eta'][m], D['Psi1'][m]
    S, N1 = eta * g, Psi1 * g * g
    Ds = np.hypot(N1, 2 * S)
    e0, c0 = g * S / Ds, Ds**2 / (g * S)                        # no solvent subtraction
    Sp = S - ETA_S * g                                           # polymer shear stress
    Dp = np.hypot(N1, 2 * Sp)
    e1, c1 = g * Sp / Dp, Dp**2 / (g * Sp) + 4 * ETA_S           # solvent subtracted, then added back (Sec. V C)
    return dict(g=g, eps=e0, P1=c0, eps_s=e1, P1_s=c1, solvent_share=ETA_S * g / S)


MODELS = {'linear PTT': (lptt_shear, lptt_ext, (1e-4, 1.0)),
          'exponential PTT': (eptt_shear, eptt_ext, (1e-4, 1.0)),
          'Giesekus': (giesekus_shear, giesekus_ext, (1e-3, 0.499))}


def multimode_fits(D, spec):
    tk, hk = np.array(spec['tk']), np.array(spec['hk'])
    st = D['gd'] < STABLE
    mN = st & (D['gd'] > 0.28)
    gd, eta, Psi1 = D['gd'], D['eta'], D['Psi1']
    out = {}
    for name, (sh, ex, bnds) in MODELS.items():
        def pk_of(q):
            return np.clip(np.exp(q[0]) * (tk / tk.max())**q[1], bnds[0] * 1e-2, bnds[1])

        def shear_curves(pk, g):
            S = np.array([sum(sh(h, t, p, x)[0] for t, h, p in zip(tk, hk, pk)) for x in g])
            N1 = np.array([sum(sh(h, t, p, x)[1] for t, h, p in zip(tk, hk, pk)) for x in g])
            return S / g + ETA_S, N1 / g**2

        def cost(q):
            e, P = shear_curves(pk_of(q), gd[st])
            return np.sum(np.log(e / eta[st])**2) + np.sum(np.log(P[mN[st]] / Psi1[mN])**2)
        best = min((minimize(cost, [np.log(p0), m0], method='Nelder-Mead', options=dict(xatol=1e-4, fatol=1e-8, maxiter=600))
                    for p0 in [0.05, 0.3] for m0 in [-0.5, 0.0, 0.5]), key=lambda r: r.fun)
        pk = pk_of(best.x)
        gg = np.logspace(-1, 1.85, 40)
        e, P = shear_curves(pk, gg)
        xx = np.logspace(-1.2, 0.85, 50)
        tr = np.array([sum(ex(h, t, p, y) for t, h, p in zip(tk, hk, pk)) for y in xx]) + 4 * ETA_S
        out[name] = dict(p1=float(np.exp(best.x[0])), m=best.x[1], rms=np.sqrt(best.fun / (st.sum() + mN.sum())),
                         g=gg, eta=e, Psi1=P, eps=xx, P1=tr)
        print(f'  {name:16s}: p1 = {np.exp(best.x[0]):.3f}, m = {best.x[1]:.2f}, rms log misfit = {out[name]["rms"]:.3f}; '
              f'eta_P1 at 1 and 5 1/s: {np.interp(1, xx, tr):.1f}, {np.interp(5, xx, tr):.1f} Pa s', flush=True)
    return out


def rolie_poly_km(D):
    """Single-mode Rolie-Poly with the parameters fitted to the flow curve in King & McKinley:
    G = 5.18 Pa, tau_d = 0.94 s, tau_R = 0.75 s, beta = 1.28, delta = 0.305, Lmax = 20."""
    rp = rolie_poly(td=0.94, tR=0.75, beta=1.28, delta=0.305, Lmax=20.0)
    G = 5.18
    gg = np.logspace(-1, 1.85, 30)
    e, P = [], []
    for x in gg:
        s = G * rp['stress'](steady_state(rp, L_shear(x), T=300))
        e.append(s[0, 1] / x)
        P.append((s[0, 0] - s[1, 1]) / x**2)
    xx = np.logspace(-1.2, 0.85, 40)
    tr = [G * (lambda s: (s[0, 0] - s[1, 1]) / x)(rp['stress'](steady_state(rp, L_ext(x), T=3000))) for x in xx]
    st = D['gd'] < STABLE
    mN = st & (D['gd'] > 0.28)
    e_d = np.interp(np.log(D['gd'][st]), np.log(gg), e)
    P_d = np.interp(np.log(D['gd'][mN]), np.log(gg), P)
    rms = np.sqrt((np.sum(np.log(e_d / D['eta'][st])**2) + np.sum(np.log(P_d / D['Psi1'][mN])**2)) / (st.sum() + mN.sum()))
    x = D['w'] * 0.94
    lve = np.sqrt(np.mean(np.concatenate([np.log(G * x * x / (1 + x * x) / D['Gp']),
                                          np.log((G * x / (1 + x * x) + ETA_S * D['w']) / D['Gpp'])])**2))
    return dict(g=gg, eta=np.array(e), Psi1=np.array(P), eps=xx, P1=np.array(tr), rms=rms, lve_rms=lve)


def compute(path):
    D = read_data(path)
    print('spectra ...', flush=True)
    est = spectra(D)
    comp = composite(D)
    print('multimode fits ...', flush=True)
    mm = multimode_fits(D, est['7 modes'])
    print('single-mode Rolie-Poly ...', flush=True)
    rp = rolie_poly_km(D)
    lo = D['w'] < 0.2
    slope_lo = np.polyfit(np.log(D['w'][:3]), np.log(D['Gp'][:3]), 1)[0]
    hi = D['w'] >= 10
    slope_hi = np.polyfit(np.log(D['w'][hi]), np.log(D['Gp'][hi]), 1)[0]
    tolist = lambda o: {k: (v.tolist() if hasattr(v, 'tolist') else v) for k, v in o.items()}
    out = dict(data={k: v.tolist() for k, v in D.items()}, spectra=est, composite=tolist(comp),
               models={k: tolist(v) for k, v in mm.items()}, rolie_poly=tolist(rp), slope_lo=slope_lo, slope_hi=slope_hi)
    save_json('fig7.json', out)
    return out


def plot(d):
    import matplotlib.ticker
    plt = setup_matplotlib()
    D = {k: np.array(v) for k, v in d['data'].items()}
    st = D['gd'] < STABLE
    fig, axs = plt.subplots(1, 3, figsize=(7.0, 2.7), constrained_layout=True)
    ax = axs[0]
    ax.loglog(D['w'], D['Gp'], 'o', mfc='none', mec=INK, ms=3.5, mew=0.7, label=r"$G'$ data")
    ax.loglog(D['w'], D['Gpp'], 's', mfc=INK, mec=INK, ms=2.8, label=r"$G''$ data")
    ww = np.logspace(-1.1, 2.1, 100)
    s = d['spectra']['7 modes']
    a, b = maxwell_moduli(np.array(s['tk']), np.array(s['hk']), ww)
    ax.loglog(ww, a, color=PALETTE[0], label='7-mode spectrum')
    ax.loglog(ww, b + ETA_S * ww, color=PALETTE[0])
    x = ww * 0.94
    ax.loglog(ww, 5.18 * x * x / (1 + x * x), color=MUTE, ls='--', label='single mode,\nKing & McKinley (2026)')
    ax.loglog(ww, 5.18 * x / (1 + x * x) + ETA_S * ww, color=MUTE, ls='--')
    ax.set_xlabel(r'$\omega$ [rad/s]')
    ax.set_ylabel(r"$G'$, $G''$ [Pa]")
    ax.set_ylim(0.02, 60)
    ax.legend(loc='lower right', fontsize=6.0, handlelength=1.6)
    panel_label(ax, '(a)')
    names = [('linear PTT', 'PTT (linear)', PALETTE[0], '-'), ('exponential PTT', 'PTT (exponential)', PALETTE[1], '--'),
             ('Giesekus', 'Giesekus', PALETTE[2], '-.')]
    ax = axs[1]
    ax.loglog(D['gd'][st], D['eta'][st], 'o', mfc='none', mec=INK, ms=3.5, mew=0.7, label=r'$\eta$ data')
    mN = st & (D['gd'] > 0.22)
    ax.loglog(D['gd'][mN], D['Psi1'][mN], 's', mfc=INK, mec=INK, ms=2.8, label=r'$\Psi_1$ data')
    for k, lb, c, ls in names:
        r = d['models'][k]
        ax.loglog(r['g'], r['eta'], color=c, ls=ls)
        ax.loglog(r['g'], r['Psi1'], color=c, ls=ls)
    r = d['rolie_poly']
    ax.loglog(r['g'], r['eta'], color=MUTE, ls='--')
    ax.loglog(r['g'], r['Psi1'], color=MUTE, ls='--')
    ax.set_xlabel(r'$\dot\gamma$ [s$^{-1}$]')
    ax.set_ylabel(r'$\eta$ [Pa s], $\Psi_1$ [Pa s$^2$]')
    ax.set_xlim(0.08, 90)
    ax.legend(loc='lower left', fontsize=6.3)
    panel_label(ax, '(b)', 'upper right')
    ax = axs[2]
    c_ = d['composite']
    ax.loglog(c_['eps_s'], c_['P1_s'], '*', mfc='none', mec=INK, ms=6.5, mew=0.8, label=r'composite $\eta_{[P1]}$ from data')
    for k, lb, c, ls in names:
        r = d['models'][k]
        ax.loglog(r['eps'], r['P1'], color=c, ls=ls, label=lb + ', multimode')
    r = d['rolie_poly']
    ax.loglog(r['eps'], r['P1'], color=MUTE, ls='--', label='Rolie-Poly, King & McKinley (2026)')
    ax.set_xlabel(r'$\dot\epsilon$, $\dot\epsilon_{\rm eff}$ [s$^{-1}$]')
    ax.set_ylabel(r'$\eta_{P1}$ [Pa s]')
    ax.set_xlim(0.08, 8)
    ax.set_ylim(10, 400)
    ax.legend(loc='upper left', fontsize=6.3)
    panel_label(ax, '(c)', 'lower right')
    ax.yaxis.set_minor_formatter(matplotlib.ticker.NullFormatter())
    fig.savefig(f'{FIGS}/fig7_pib.pdf')
    plt.close(fig)


def check(d):
    ck = Checks('fig7')
    S = d['spectra']
    nn = [S[k] for k in ('5 modes', '7 modes', '7 modes, to 30 s', '9 modes')]
    allsp = nn + [S['chain']]
    ck.add('chain-like spectrum exponent k', '1.90', S['chain']['k'], tol=0.006)
    ck.add('uncertainty of k', '0.02', S['chain']['dk'], tol=0.006)
    ck.add("slope of G' above 10 rad/s", '0.51', d['slope_hi'])
    D = {k: np.array(v) for k, v in d['data'].items()}
    sl = lambda G, a, b: np.polyfit(np.log(D['w'][a:b]), np.log(D[G][a:b]), 1)[0]
    ck.add("slope of G'' between 0.1 and 0.4 rad/s", '0.8', sl('Gpp', 0, 7))
    ck.add("slope of G' between 0.1 and 0.4 rad/s, without the lowest point", '1.4', sl('Gp', 1, 7))
    ck.add("slope of G' between 0.1 and 0.4 rad/s, all points", '1.6', sl('Gp', 0, 7))
    ck.add('NNLS rms relative residuals: smallest', '0.05', min(s['rms'] for s in nn))
    ck.add('NNLS rms relative residuals: largest', '0.07', max(s['rms'] for s in nn))
    for s_ in allsp:
        ck.add('eta0 of each spectrum', '3.8', s_['eta0'])
    ck.add('<tau>: smallest over spectra', '1.3', min(s_['t1'] for s_ in allsp))
    ck.add('<tau>: largest over spectra', '1.7', max(s_['t1'] for s_ in allsp))
    ck.add('<tau>^2/<tau^2>: smallest', '0.3', min(s_['t1']**2 / s_['t2'] for s_ in allsp))
    ck.add('<tau>^2/<tau^2>: largest', '0.7', max(s_['t1']**2 / s_['t2'] for s_ in allsp))
    lead = [16 * (s_['t2'] - s_['t1']**2) * 0.01 / 4 for s_ in allsp]          # 16 Var eps^2 / 4, eps = 0.1 1/s
    ck.add('leading correction at 0.1 1/s as a fraction of 4 eta0: smallest (3%)', '0.03', min(lead))
    ck.add('leading correction at 0.1 1/s as a fraction of 4 eta0: largest (27%)', '0.27', max(lead))
    ck.add('Rouse-like coil-stretch rate 1/2tau_1 from the chain spectrum, low end', '0.25', 0.5 / S['chain']['tau1'], tol=0.03)
    c = {k: np.array(v) for k, v in d['composite'].items()}
    ck.add('solvent share of the shear stress at 63 1/s (3%)', '0.03', c['solvent_share'][-1])
    ck.add('largest change of eta_[P1] by solvent subtraction (at most 3%)', '0.03', np.max(abs(c['P1_s'] / c['P1'] - 1)), tol=0.005)
    ck.add('4 eta0 [Pa s]', '15.3', 4 * S['7 modes']['eta0'], tol=0.25)
    sel = c['eps_s'] <= 1.0
    ck.add('composite / 4 eta0 for eps_eff <= 1: smallest (10% above)', '1.10', (c['P1_s'][sel] / (4 * S['7 modes']['eta0'])).min(), tol=0.03)
    ck.add('composite / 4 eta0 for eps_eff <= 1: largest (25% above)', '1.25', (c['P1_s'][sel] / (4 * S['7 modes']['eta0'])).max(), tol=0.03)
    ck.add('composite at the highest rate [Pa s]', '70', c['P1_s'][-1], tol=1.5)
    ck.add('highest eps_eff [1/s]', '5.7', c['eps_s'][-1])
    interp = lambda x: np.exp(np.interp(np.log(x), np.log(c['eps_s']), np.log(c['P1_s'])))
    ck.add('composite at 1 1/s [Pa s]', '19', interp(1.0))
    ck.add('composite at 5 1/s [Pa s]', '55', interp(5.0))
    M = d['models']
    for k, q1, q5 in [('linear PTT', '35', '42'), ('exponential PTT', '24', '26'), ('Giesekus', '35', '56')]:
        ck.add(f'{k}: eta_P1 at 1 1/s', q1, np.interp(1.0, M[k]['eps'], M[k]['P1']))
        ck.add(f'{k}: eta_P1 at 5 1/s', q5, np.interp(5.0, M[k]['eps'], M[k]['P1']))
    ck.add('multimode fits: smallest rms log misfit', '0.14', min(M[k]['rms'] for k in M))
    ck.add('multimode fits: largest rms log misfit', '0.21', max(M[k]['rms'] for k in M))
    ck.add('single-mode Rolie-Poly: rms log misfit', '0.21', d['rolie_poly']['rms'])
    ck.add('single mode vs SAOS: rms factor', '3.5', np.exp(d['rolie_poly']['lve_rms']))
    rp = d['rolie_poly']
    dev = np.max(abs(np.interp(np.log(c['eps_s']), np.log(rp['eps']), np.log(rp['P1'])) - np.log(c['P1_s'])))
    ck.at_most('composite within 18% of the single-mode Rolie-Poly eta_P1', '0.18', np.exp(dev) - 1)
    return ck.save()


if __name__ == '__main__':
    if '--plot' in sys.argv:
        d = load_json('fig7.json')
    else:
        if len(sys.argv) < 2 or not os.path.exists(sys.argv[1]):
            sys.exit(__doc__)
        d = compute(sys.argv[1])
    plot(d)
    check(d)
